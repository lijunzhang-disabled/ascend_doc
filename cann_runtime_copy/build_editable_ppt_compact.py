"""Wider, shorter editable composition for the right half of a slide."""
from build_editable_ppt import Drawing, NOTES, OUT
from pptx import Presentation
from pptx.util import Inches
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN


def draw(slide, scale=1., ox=0., oy=0.):
    d = Drawing(slide, scale, ox, oy)
    d.text(.25, .12, 8.70, .43, '加密拷贝：同步与异步', 29, bold=True)
    d.text(.25, .61, 8.70, .26, 'CANN Runtime · Ascend 950 / UB · runtime_enc / iter-014', 15, bold=True)

    for x, color, title, api, body in [
        (.25, 'blue', '同步接口', 'aclrtMemcpy', '整条加密链完成后返回'),
        (4.72, 'teal', '异步接口', 'aclrtMemcpyAsync', '提交返回；流同步等待完成'),
    ]:
        d.box(x, 1.02, 4.23, .92, color+'_bg', color, name=title+'卡片')
        d.box(x, 1.02, 4.23, .29, color, name=title+'标题底色')
        d.text(x+.06, 1.025, 4.11, .27, title, 17.5, 'white', True)
        d.text(x+.06, 1.37, 4.11, .25, api, 17, bold=True)
        d.text(x+.06, 1.68, 4.11, .21, body, 14)
    d.line(2.365, 1.95, 2.365, 2.03, arrow=False)
    d.line(6.835, 1.95, 6.835, 2.03, arrow=False)
    d.line(2.365, 2.03, 6.835, 2.03, arrow=False)
    d.line(4.60, 2.03, 4.60, 2.14)
    d.box(.25, 2.16, 8.70, .36, 'summary', 'blue', name='共同安全执行路径')
    d.richrow(.35, 2.18, 8.50, .29, [
        ('共同进入 SecureDma', 'navy', True),
        ('：Host worker 负责主机处理与密文搬运', 'navy', False),
    ], 16)

    d.box(.25, 2.67, 8.70, 1.06, 'teal_bg', 'teal', name='设备控制路径')
    d.box(.25, 2.67, 8.70, .30, 'teal', name='控制路径标题底色')
    d.text(.35, 2.675, 8.50, .28, '设备流上的控制路径', 17.5, 'white', True)
    nodes = [
        (.43, 2.02, 'Runtime 提交 SQE', 'navy', 'white'),
        (2.68, 1.84, 'STARS 按流调度', 'navy', 'white'),
        (4.75, 1.90, 'ValueWait(READY)', 'purple', 'purple_bg'),
        (6.88, 1.89, 'AICPU SecureDma', 'purple', 'purple_bg'),
    ]
    for x, w, text, color, bg in nodes:
        d.box(x, 3.10, w, .34, bg, 'teal' if bg=='white' else 'purple', name=text+'节点')
        d.text(x+.02, 3.10, w-.04, .34, text, 14.5, color, True)
    for left, right in zip(nodes, nodes[1:]):
        d.line(left[0]+left[1]+.015, 3.27, right[0]-.025, 3.27, 'teal')
    d.text(.40, 3.51, 8.40, .17, 'STARS 调度任务，AICPU 执行加解密', 12.8)

    def section(y, color, title, height):
        d.box(.25, y, 8.70, height, 'white', color, name=title+'面板')
        d.box(.25, y, 8.70, .31, color, name=title+'标题底色')
        d.text(.43, y+.005, 8.34, .285, title, 17.5, 'white', True, PP_ALIGN.LEFT)

    def step(y, n, role, color, parts, size=15.5):
        bg = 'amber_bg' if color=='amber' else color+'_bg'
        d.box(.43, y, 1.40, .31, bg, name=role+'步骤标签')
        d.text(.47, y, .24, .31, str(n), 15, color, True)
        d.text(.73, y, 1.04, .31, role, 14.5, color, True)
        d.box(1.99, y, 6.78, .31, bg, 'line', name=role+'流程背景')
        d.richrow(2.03, y, 6.70, .31, parts, size)

    section(3.88, 'blue', 'H2D：Host 加密，Device 解密', 1.76)
    step(4.29, 1, 'Host 侧', 'blue', [
        ('业务明文 → ', 'navy', False), ('Host worker 加密', 'purple', True), (' → 密文缓冲', 'navy', False)])
    d.line(5.38, 4.61, 5.38, 4.71)
    step(4.72, 2, 'UB 传输', 'amber', [
        ('驱动 MemCopySync', 'amber', True), (' → URMA → Device 密文缓冲', 'navy', False)], 15.2)
    d.text(2.06, 5.07, 6.05, .17, '密文搬运完成后发布 READY', 12, 'muted', align=PP_ALIGN.LEFT)
    d.line(8.40, 5.05, 8.40, 5.26, 'blue', True, 1.2)
    step(5.27, 3, 'Device 侧', 'teal', [
        ('ValueWait 放行 → ', 'navy', False), ('AICPU 解密', 'purple', True), (' → 业务明文', 'navy', False)])

    section(5.80, 'teal', 'D2H：Device 加密，Host 解密', 1.60)
    step(6.20, 1, 'Device 侧', 'teal', [
        ('业务明文 → ', 'navy', False), ('AICPU 加密', 'purple', True), (' → 密文缓冲', 'navy', False)])
    d.line(5.38, 6.52, 5.38, 6.60, 'teal')
    step(6.61, 2, 'UB 传输', 'amber', [
        ('设备加密完成 → ', 'navy', False), ('驱动 / URMA', 'amber', True), (' → Host 密文缓冲', 'navy', False)], 15.2)
    d.line(5.38, 6.93, 5.38, 7.01, 'teal')
    step(7.02, 3, 'Host 侧', 'blue', [
        ('Host worker 解密', 'purple', True), (' → 业务明文', 'navy', False)])

    d.box(.25, 7.55, 8.70, .62, 'summary', 'blue', name='完成语义总结')
    d.text(.40, 7.58, 8.40, .26, '链路搬运密文；端侧处理明文', 17.5, bold=True)
    d.text(.40, 7.88, 8.40, .26, '完成语义覆盖：加密 + 搬运 + 解密', 15.5, bold=True)
    d.text(.30, 8.26, 8.60, .20,
           'AES-256-CTR · 小拷贝整帧 / 大拷贝按片流水 · 加密路径启用时', 11.5, 'muted')
    return d


def build(name, width, height, scale, ox, oy):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(width), Inches(height)
    prs.core_properties.title = 'CANN Runtime 加密拷贝：同步与异步（宽版）'
    prs.core_properties.subject = 'Ascend 950 UB / iter-014，原生可编辑流程图'
    prs.core_properties.author = 'Codex'
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
    d = draw(slide, scale, ox, oy)
    slide.notes_slide.notes_text_frame.text = NOTES
    dest = OUT / name
    prs.save(dest)
    print(dest, '| native elements=', len(d.shapes))


if __name__ == '__main__':
    build('cann_secure_copy_editable_right_panel_wide.pptx', 13.333333, 7.5, .69, 6.826667, .78)
    build('cann_secure_copy_editable_asset_wide.pptx', 9.2, 8.6, 1., 0., 0.)
