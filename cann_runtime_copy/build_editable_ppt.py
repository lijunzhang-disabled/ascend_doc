"""Build editable native PowerPoint diagrams for the verified secure-copy path."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement

OUT = Path(__file__).resolve().parent
FONT = 'Microsoft YaHei'
C = {
    'navy': '122448', 'blue': '0879D9', 'teal': '009BA3',
    'purple': '713BC0', 'amber': 'B96D05', 'muted': '63748A',
    'white': 'FFFFFF', 'blue_bg': 'EEF6FE', 'teal_bg': 'ECFAF8',
    'purple_bg': 'F4EDFE', 'amber_bg': 'FFF5DF', 'line': 'B4CFE5',
    'summary': 'EAF4FD',
}


class Drawing:
    def __init__(self, slide, scale=1., ox=0., oy=0.):
        self.scale, self.ox, self.oy = scale, ox, oy
        self.root = slide.shapes.add_group_shape()
        self.root.name = '加密拷贝流程图（取消组合后逐项编辑）'
        self.shapes = self.root.shapes
        self.text_items = []

    def x(self, x): return Inches(self.ox + x*self.scale)
    def y(self, y): return Inches(self.oy + y*self.scale)
    def n(self, n): return Inches(n*self.scale)

    def box(self, x, y, w, h, fill='white', border=None, radius=True, name='框'):
        sh = self.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
            self.x(x), self.y(y), self.n(w), self.n(h))
        sh.name = name
        sh._element.spPr.append(OxmlElement('a:effectLst'))
        if radius:
            sh.adjustments[0] = 0.12
        sh.fill.solid()
        sh.fill.fore_color.rgb = RGBColor.from_string(C.get(fill, fill))
        if border:
            sh.line.color.rgb = RGBColor.from_string(C.get(border, border))
            sh.line.width = Pt(.8*self.scale)
        else:
            sh.line.fill.background()
        return sh

    def text(self, x, y, w, h, text, size=16, color='navy', bold=False,
             align=PP_ALIGN.CENTER, name=None):
        sh = self.shapes.add_textbox(self.x(x), self.y(y), self.n(w), self.n(h))
        sh.name = name or text.replace('\n', ' / ')
        tf = sh.text_frame
        tf.clear()
        tf.word_wrap = False
        tf.margin_left = tf.margin_right = Inches(.02*self.scale)
        tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        for i, line in enumerate(text.split('\n')):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            p.space_before = p.space_after = Pt(0)
            p.line_spacing = 1.08
            r = p.add_run()
            r.text = line
            r.font.name = FONT
            r.font.size = Pt(size*self.scale)
            r.font.bold = bold
            r.font.color.rgb = RGBColor.from_string(C.get(color, color))
            rpr = r._r.get_or_add_rPr()
            rpr.set('lang', 'zh-CN')
            ea = OxmlElement('a:ea')
            ea.set('typeface', FONT)
            rpr.append(ea)
        self.text_items.append((text, x, y, w, h, size))
        return sh

    def line(self, x1, y1, x2, y2, color='blue', arrow=True, width=1.6, name='连接箭头'):
        sh = self.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, self.x(x1), self.y(y1), self.x(x2), self.y(y2))
        sh.name = name
        sh.line.color.rgb = RGBColor.from_string(C[color])
        sh.line.width = Pt(width*self.scale)
        if arrow:
            end = OxmlElement('a:tailEnd')
            end.set('type', 'triangle')
            end.set('w', 'med')
            end.set('len', 'med')
            sh.line._get_or_add_ln().append(end)
        return sh

    def richrow(self, x, y, w, h, parts, size=15.5):
        sh = self.text(x, y, w, h, '', size)
        sh.name = ''.join(t for t, _, _ in parts)
        p = sh.text_frame.paragraphs[0]
        for t, color, bold in parts:
            r = p.add_run()
            r.text = t
            r.font.name = FONT
            r.font.size = Pt(size*self.scale)
            r.font.bold = bold
            r.font.color.rgb = RGBColor.from_string(C[color])
            rpr = r._r.get_or_add_rPr()
            rpr.set('lang', 'zh-CN')
            ea = OxmlElement('a:ea')
            ea.set('typeface', FONT)
            rpr.append(ea)
        return sh


def draw(slide, scale=1., ox=0., oy=0.):
    d = Drawing(slide, scale, ox, oy)
    d.text(.25, .12, 7.5, .48, '加密拷贝：同步与异步', 29, bold=True)
    d.text(.25, .64, 7.5, .29, 'CANN Runtime · Ascend 950 / UB', 17, bold=True)
    d.box(1.35, 1.00, 5.3, .30, 'blue_bg', name='实现版本')
    d.text(1.37, 1.00, 5.26, .30, 'runtime_enc · iter-014 · 加密路径启用时', 13)

    # API contracts. Both converge on the same encrypted implementation.
    for x, color, title, api, lines in [
        (.25, 'blue', '同步接口', 'aclrtMemcpy', '整条加密链完成后返回'),
        (4.10, 'teal', '异步接口', 'aclrtMemcpyAsync', '提交返回；流同步等待完成'),
    ]:
        d.box(x, 1.44, 3.65, 1.17, color+'_bg', color, name=title+'卡片')
        d.box(x, 1.44, 3.65, .34, color, name=title+'标题底色')
        d.text(x+.08, 1.45, 3.49, .31, title, 18, 'white', True)
        d.text(x+.08, 1.86, 3.49, .28, api, 17, bold=True)
        d.text(x+.07, 2.19, 3.51, .29, lines, 14.0)
    d.line(2.08, 2.61, 2.08, 2.74, arrow=False)
    d.line(5.92, 2.61, 5.92, 2.74, arrow=False)
    d.line(2.08, 2.74, 5.92, 2.74, arrow=False)
    d.line(4, 2.74, 4, 2.88)
    d.box(.25, 2.90, 7.5, .57, 'summary', 'blue', name='共同安全执行路径')
    d.text(.35, 2.93, 7.3, .26, '共同进入 SecureDma', 20, bold=True)
    d.text(.35, 3.20, 7.3, .20, 'Host worker 负责主机处理与密文搬运', 14)

    # STARS control flow is distinct from ciphertext transport.
    d.box(.25, 3.61, 7.5, 1.52, 'teal_bg', 'teal', name='设备控制路径')
    d.box(.25, 3.61, 7.5, .34, 'teal', name='控制路径标题底色')
    d.text(.35, 3.62, 7.3, .31, '设备流上的控制路径', 18, 'white', True)
    for x, y, w, text, color, bg in [
        (1.00, 4.05, 2.70, 'Runtime 提交 SQE', 'navy', 'white'),
        (4.30, 4.05, 2.70, 'STARS 按流调度', 'navy', 'white'),
        (1.00, 4.54, 2.70, 'ValueWait(READY)', 'purple', 'purple_bg'),
        (4.30, 4.54, 2.70, 'AICPU SecureDma', 'purple', 'purple_bg'),
    ]:
        d.box(x, y, w, .32, bg, 'teal' if bg=='white' else 'purple', name=text+'节点')
        d.text(x+.03, y, w-.06, .32, text, 16, color, True)
    d.line(3.72, 4.21, 4.24, 4.21, 'teal')
    # Elbowed sequence: STARS -> ValueWait -> AICPU; no bypass of the gate.
    d.line(5.65, 4.38, 5.65, 4.45, 'teal', False, 1.2)
    d.line(5.65, 4.45, 2.35, 4.45, 'teal', False, 1.2)
    d.line(2.35, 4.45, 2.35, 4.53, 'teal', True, 1.2)
    d.line(3.72, 4.70, 4.24, 4.70, 'teal')
    d.text(.40, 4.91, 7.20, .17, 'STARS 调度任务，AICPU 执行加解密', 12.8)

    def section(y, color, title, height):
        d.box(.25, y, 7.50, height, 'white', color, name=title+'面板')
        d.box(.25, y, 7.50, .36, color, name=title+'标题底色')
        d.text(.40, y+.01, 7.20, .32, title, 18, 'white', True, PP_ALIGN.LEFT)

    def step(y, num, role, color, parts, size=15.2):
        bg = 'amber_bg' if color=='amber' else color+'_bg'
        d.box(.40, y, 1.36, .34, bg, name=role+'步骤标签')
        d.text(.44, y, .24, .34, str(num), 15, color, True)
        d.text(.69, y, 1.00, .34, role, 14.5, color, True)
        d.box(1.90, y, 5.65, .34, bg, 'line', name=role+'流程背景')
        d.richrow(1.94, y, 5.57, .34, parts, size)

    section(5.28, 'blue', 'H2D：Host 加密，Device 解密', 1.89)
    step(5.72, 1, 'Host 侧', 'blue', [
        ('明文 → ', 'navy', False), ('Host worker 加密', 'purple', True), (' → 密文缓冲', 'navy', False)])
    d.line(4.73, 6.07, 4.73, 6.17)
    step(6.18, 2, 'UB 传输', 'amber', [
        ('MemCopySync', 'amber', True), (' → URMA → Device 密文', 'navy', False)], 15.0)
    d.text(1.94, 6.55, 4.9, .18, '密文搬运完成后发布 READY', 12, 'muted', align=PP_ALIGN.LEFT)
    d.line(7.15, 6.54, 7.15, 6.74, 'blue', True, 1.2)
    step(6.75, 3, 'Device 侧', 'teal', [
        ('ValueWait → ', 'navy', False), ('AICPU 解密', 'purple', True), (' → 明文', 'navy', False)])

    section(7.29, 'teal', 'D2H：Device 加密，Host 解密', 1.75)
    step(7.73, 1, 'Device 侧', 'teal', [
        ('明文 → ', 'navy', False), ('AICPU 加密', 'purple', True), (' → 密文缓冲', 'navy', False)])
    d.line(4.73, 8.08, 4.73, 8.18, 'teal')
    step(8.19, 2, 'UB 传输', 'amber', [
        ('加密完成 → ', 'navy', False), ('驱动 / URMA', 'amber', True), (' → Host 密文', 'navy', False)], 14.8)
    d.line(4.73, 8.54, 4.73, 8.64, 'teal')
    step(8.65, 3, 'Host 侧', 'blue', [
        ('Host worker 解密', 'purple', True), (' → 业务明文', 'navy', False)])

    d.box(.25, 9.16, 7.5, .63, 'summary', 'blue', name='完成语义总结')
    d.text(.40, 9.21, 7.2, .25, '链路搬运密文；端侧处理明文', 19, bold=True)
    d.text(.40, 9.50, 7.2, .23, '完成语义覆盖：加密 + 搬运 + 解密', 16.5, bold=True)
    d.text(.35, 9.87, 7.30, .19, '当前算法：AES-256-CTR', 12.5, 'muted')
    d.text(.35, 10.07, 7.30, .18, '小拷贝整帧处理；大拷贝按片流水执行。', 12, 'muted')
    return d


NOTES = '''图中所有文字、方框、流程箭头均为原生 PowerPoint 元素，可取消组合后编辑。

技术范围：Ascend 950 / UB，runtime_enc 的 iter-014，加密路径启用时。
Runtime HEAD: 9b9efba50；AICPU HEAD: 7fd3986。
同步和异步都路由到 SecureDma；同步等待完整处理结束，异步提交后返回。
密文搬运由 Host worker 通过 Driver::MemCopySync / UB 驱动 URMA 发起。
STARS 调度 ValueWait 和 AICPU SecureDma；AICPU 负责设备侧加解密。
H2D：Host 加密 -> 密文搬运 -> 发布 READY -> AICPU 解密。
D2H：AICPU 加密 -> 完成 -> 密文搬运 -> Host 解密。
Stream 同步通过 HostCompletion 纳入主机解密完成。
算法是 AES-256-CTR，本图未声明 AEAD 或 MAC。

源代码：
/home/zlj/ub/runtime_enc/AscendCCv2-Runtime/src/runtime/core/src/memory/secure_dma/
  secure_dma_engine.cc
  secure_dma_whole_copy.cc
  secure_dma_large_copy.cc
  secure_dma_memory.cc
  secure_dma_runtime.cc
  secure_dma_aes_ctr.cc
/home/zlj/ub/runtime_enc/AscendCCv2-Runtime/src/runtime/core/src/stream/stream.cc
/home/zlj/ub/runtime_enc/AscendCCv2-AICPU/ms_kernels/src/secure_dma/secure_dma_whole_copy_device.inc
'''


def build(name, width, height, scale, ox, oy):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(width), Inches(height)
    prs.core_properties.title = 'CANN Runtime 加密拷贝：同步与异步'
    prs.core_properties.subject = 'Ascend 950 UB / runtime_enc iter-014，可编辑流程图'
    prs.core_properties.author = 'Codex'
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
    drawing = draw(slide, scale, ox, oy)
    slide.notes_slide.notes_text_frame.text = NOTES
    dest = OUT / name
    prs.save(dest)
    return dest, drawing


if __name__ == '__main__':
    for args in [
        ('cann_secure_copy_editable_right_panel.pptx', 13.333333, 7.5, .68, 7.28, .20),
        ('cann_secure_copy_editable_portrait.pptx', 8.0, 10.40, 1.0, .0, .0),
    ]:
        path, drawing = build(*args)
        print(f'{path} | native elements={len(drawing.shapes)}')
