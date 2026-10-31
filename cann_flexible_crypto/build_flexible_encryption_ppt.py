"""One editable slide: hardware channel protection vs proposed data-aware policy."""
from pathlib import Path
import sys
from pptx import Presentation
from pptx.util import Inches
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / 'cann_runtime_copy'))
from build_editable_ppt import Drawing, C

OUT = Path(__file__).resolve().parent
C.update({
    'nv': '5A8E12', 'nv_bg': 'F1F7E9', 'nv_line': 'CFDFBA',
    'navy': '152744', 'blue': '1C6DD0', 'blue_bg': 'EEF5FF',
    'purple': '7040B8', 'purple_bg': 'F2ECFA',
    'teal': '008C91', 'teal_bg': 'EAF7F6',
    'panel': 'FAFCFF', 'line': 'DFE6EE', 'muted': '607087',
})

WHITEPAPER = 'https://docs.nvidia.com/nvidia-secure-ai-with-blackwell-and-hopper-gpus-whitepaper.pdf'
GPU_MODE = 'https://docs.nvidia.com/confidential-computing/virtual-machines/dev/deployment-guide/gpu-mode'


class FlatDrawing(Drawing):
    def box(self, *args, **kwargs):
        sh = super().box(*args, **kwargs)
        style = sh._element.find(qn('p:style'))
        if style is not None:
            effect = style.find(qn('a:effectRef'))
            if effect is not None:
                effect.set('idx', '0')
        return sh


def create():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333333), Inches(7.5)
    prs.core_properties.title = 'CANN Runtime 可选加密：按业务流量选择保护路径'
    prs.core_properties.subject = 'NVIDIA 硬件 Inline 与 Runtime 按数据类别分流的方案对比'
    prs.core_properties.author = 'Codex'
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
    d = FlatDrawing(slide)
    d.root.name = '完整方案对比页（原生可编辑元素）'

    # Headline.
    d.text(.46, .25, 10.8, .48, 'CANN Runtime 可选加密', 29, bold=True, align=PP_ALIGN.LEFT)
    d.text(.48, .83, 12.28, .27,
           '按数据类别选择保护路径：模型权重与常规 H2D/D2H 加密，KV cache 可按策略旁路',
           15.2, 'muted', align=PP_ALIGN.LEFT)
    d.box(11.66, .35, 1.20, .33, 'blue_bg', name='方案状态')
    d.text(11.68, .35, 1.16, .33, '方案构想', 12.5, 'blue', True)

    # Panels use equally strong geometry; the comparison is policy granularity.
    d.box(.45, 1.39, 5.94, 4.91, 'panel', 'line', name='NVIDIA 对比面板')
    d.box(6.73, 1.39, 6.15, 4.91, 'panel', 'line', name='CANN 方案面板')
    d.box(.45, 1.39, 5.94, .055, 'nv', radius=False, name='NVIDIA 顶部色条')
    d.box(6.73, 1.39, 6.15, .055, 'blue', radius=False, name='CANN 顶部色条')
    d.text(.70, 1.62, 5.43, .34, 'NVIDIA · 硬件 Inline 加解密', 20.5, bold=True, align=PP_ALIGN.LEFT)
    d.text(.71, 2.02, 5.40, .24, '受保护通道内统一加密', 15, 'nv', True, PP_ALIGN.LEFT)
    d.text(6.98, 1.62, 5.64, .34, '我们的方案 · Runtime 可选加密', 20.5, bold=True, align=PP_ALIGN.LEFT)
    d.text(6.99, 2.02, 5.61, .24, '按业务标签选择加密或旁路', 15, 'blue', True, PP_ALIGN.LEFT)

    # Left: one protected channel applies the same link policy to each class.
    d.text(.72, 2.41, 1.53, .20, '待传输数据', 12, 'muted')
    d.box(2.65, 2.36, 1.72, .29, 'nv_bg', name='统一保护启用标记')
    d.text(2.67, 2.36, 1.68, .29, '保护启用：ON', 12, 'nv', True)
    d.text(4.76, 2.41, 1.36, .20, '受保护通道', 12, 'muted')
    d.box(2.66, 2.81, 1.71, 1.91, 'nv_bg', 'nv', name='硬件 inline 引擎')
    d.text(2.75, 3.10, 1.53, .32, 'Inline', 23, 'nv', True)
    d.text(2.75, 3.56, 1.53, .34, '加 / 解密', 20, 'navy', True)
    d.text(2.74, 4.12, 1.55, .23, '硬件透明执行', 11.8, 'nv')
    for idx, (y, label) in enumerate([(2.89, '模型权重'), (3.53, '常规 H2D/D2H'), (4.17, 'KV cache')], 1):
        d.box(.71, y, 1.55, .45, 'white', 'line', name=f'NVIDIA 输入 {label}')
        d.text(.73, y, 1.51, .45, label, 13.5 if idx==2 else 15, bold=True)
        d.line(2.30, y+.225, 2.60, y+.225, 'nv', width=1.8)
        d.line(4.41, y+.225, 4.70, y+.225, 'nv', width=1.8)
        d.box(4.75, y, 1.35, .45, 'nv_bg', 'nv_line', name=f'NVIDIA 输出 {label} 密文')
        d.text(4.77, y, 1.31, .45, '密文传输', 14, 'nv', True)
    d.text(.72, 4.92, 5.38, .25, '同一受保护通道内，采用统一的链路保护策略', 13.2)
    d.box(.71, 5.36, 5.39, .51, 'nv_bg', name='硬件保护粒度')
    d.text(.85, 5.43, 1.05, .30, '选择粒度', 13, 'nv', True)
    d.text(2.00, 5.43, 3.91, .30, '设备 / 受保护通道', 18, 'navy', True)
    d.text(.73, 6.00, 5.35, .20, '对应用透明 · 链路策略不区分权重与 KV 标签', 12.8, 'muted')

    # Right: explicit application semantics precede the Runtime routing choice.
    d.box(7.01, 2.43, 5.59, .36, 'blue_bg', name='业务分类入口')
    d.text(7.09, 2.43, 5.43, .36, '业务 / 框架标记：Buffer · Tensor · 内存池', 13.5, 'blue', True)
    d.line(9.805, 2.82, 9.805, 2.97, 'blue')
    d.box(7.01, 3.01, 5.59, 1.52, 'white', 'blue', name='Runtime 策略选择器')
    d.text(7.16, 3.08, 5.29, .29, 'CANN Runtime · 策略选择器', 17, 'blue', True)
    rows = [
        (3.48, '模型权重', '加密', 'purple', 'purple_bg'),
        (3.82, '常规 H2D / D2H', '加密', 'purple', 'purple_bg'),
        (4.16, 'KV cache', '旁路 / 不加密', 'teal', 'teal_bg'),
    ]
    for y, label, action, color, bg in rows:
        d.box(7.20, y, 5.21, .27, bg, name=f'策略：{label} {action}')
        d.text(7.32, y, 2.67, .27, label, 14, bold=True, align=PP_ALIGN.LEFT)
        d.text(10.30, y, 1.91, .27, action, 14, color, True)
    d.line(9.805, 4.55, 9.805, 4.71, 'blue', arrow=False)
    d.line(8.34, 4.71, 11.27, 4.71, 'blue', arrow=False)
    d.line(8.34, 4.71, 8.34, 4.87, 'purple')
    d.line(11.27, 4.71, 11.27, 4.87, 'teal')
    d.box(7.01, 4.92, 2.66, .87, 'purple_bg', 'purple', name='选择加密路径')
    d.text(7.13, 5.02, 2.42, .29, '加密路径', 18, 'purple', True)
    d.text(7.10, 5.40, 2.48, .23, '复用 SecureDma', 13.5)
    d.box(9.94, 4.92, 2.66, .87, 'teal_bg', 'teal', name='选择原生拷贝路径')
    d.text(10.04, 5.02, 2.46, .29, '旁路路径', 18, 'teal', True)
    d.text(10.04, 5.40, 2.46, .23, '复用原生拷贝', 13.5)
    d.text(7.05, 6.00, 5.51, .20, '显式标签优先；未标记流量默认加密', 13.2, 'muted')

    # Outcome and concise scope, without making unmeasured performance claims.
    d.box(.45, 6.49, 12.43, .51, 'navy', name='方案目标')
    d.text(.65, 6.56, 12.03, .32, '目标：保护策略可配置，加解密成本按需分配', 20.5, 'white', True)
    d.text(.49, 7.10, 12.34, .16,
           '注：左图限定支持 TDISP/IDE 的端到端平台与受保护通道；右图为拟议方案，KV 旁路需业务明确允许。',
           9.2, 'muted', align=PP_ALIGN.LEFT)
    s1 = d.text(.49, 7.30, 5.6, .13, '参考 [1] NVIDIA Secure AI White Paper · TDISP/IDE', 8.4, 'muted', align=PP_ALIGN.LEFT)
    s1.text_frame.paragraphs[0].runs[0].hyperlink.address = WHITEPAPER
    s2 = d.text(6.85, 7.30, 5.98, .13, '[2] NVIDIA GPU Mode · 源码基线 runtime_enc / iter-014', 8.4, 'muted', align=PP_ALIGN.RIGHT)
    s2.text_frame.paragraphs[0].runs[0].hyperlink.address = GPU_MODE

    slide.notes_slide.notes_text_frame.text = f'''展示目的
比较硬件/通道级统一保护与业务语义驱动的 Runtime 可选加密。右侧是新方案构想，未声明已经实现。

左侧来源与范围
[1] {WHITEPAPER}
NVIDIA Secure AI with Blackwell and Hopper GPUs, v1.3, August 2025.
Printed page 9 describes Direct Access with Inline Encryption Using TDISP/IDE and requires compatible CPU, GPU and (if present) IDE-capable PCIe switch.
Printed pages 13-14 distinguish configurations. Hopper PPCIE can leave NVLink unencrypted. Blackwell multi-GPU configurations also support NVLink protection. Page 21 describes hardware inline CPU-GPU encryption as platform/enablement dependent.
[2] {GPU_MODE}
GPU CC/PPCIE modes are configured per GPU by a privileged host operation. Modes vary by platform.
Accessed 2026-10-31.

不能把“NV 开就全开”表述为 NVIDIA 所有设备、所有互连、所有模式一律全量加密。本页使用受保护通道内统一链路策略来表达比较维度。“链路策略不区分权重与 KV 标签”是链路机制与应用语义分层的概括，不是穷举证明所有 NVIDIA API 都不支持任何选择能力。
同一受保护通道内的明文端点数据经 inline 加解密后在通道中密文传输。图中没有将 GPU 内部数据误画成永远密文。

右侧方案
框架/应用显式传递数据用途，通过 Buffer/Tensor/内存池标签或注册信息在 Runtime 做策略选择。普通 memcpy 参数只有地址、长度、方向与流，不足以独立判断权重或 KV。
示例规则：模型权重 ENCRYPT；常规 H2D/D2H DEFAULT=ENCRYPT；显式 KV_CACHE=BYPASS。
显式类别优先于通用方向规则；未标记保持默认加密。KV 不加密是业务策略示例，不表示 KV 天然不敏感或与全流量保护等价。
加密分支复用 SecureDma；旁路分支复用原生 Memcpy/MemcpyAsync 及其同步/异步语义。
实现仍需定义标签生命周期、内存复用后的清理、图捕获/重放时的策略稳定性及覆盖的接口范围，避免同一缓冲被误分类。性能收益有待测试，本页只陈述目标。

现有实现基线（本地只读核对）
/home/zlj/ub/runtime_enc/AscendCCv2-Runtime/src/runtime/core/src/memory/secure_dma/secure_dma_route.cc
ResolveRequest 根据 H2D/D2H 等方向和已启用 Engine 路由；当前未提供本页提出的业务类别策略。
建议的策略判断位于安全路由入口：匹配 ENCRYPT 后走既有 Engine，明确 BYPASS 返回原生拷贝路径。
本次仅制作 PPT，不修改 Runtime。

可编辑性
本页由 PowerPoint 原生文本框、形状和箭头构成，已整体组合。取消组合即可逐项编辑。
'''
    dst = OUT / 'cann_runtime_flexible_encryption.pptx'
    prs.save(dst)
    print(f'{dst}: 1 slide, {len(d.shapes)} native elements')


if __name__ == '__main__':
    create()
